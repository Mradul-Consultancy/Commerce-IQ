import argparse
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.append(str(ROOT))
sys.path.append(str(ROOT / "src"))

import analytics as a          # noqa: E402
import validate_data           # noqa: E402

REPORTS_DIR = ROOT / "reports"

MONEY_COLS = {"revenue", "profit", "lifetime_value", "avg_order_value",
              "avg_spend_per_customer", "prev_month_revenue"}
PCT_COLS = {"growth_pct", "margin_pct", "pct_of_revenue", "pass_rate_pct"}


def style_sheet(ws, df):
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        width = max([len(str(col))] + [len(f"{v:,.0f}" if isinstance(v, float) else str(v))
                                       for v in df[col].head(200)]) + 3
        ws.column_dimensions[letter].width = min(max(width, 12), 45)
        fmt = None
        if col in MONEY_COLS:
            fmt = "#,##0"
        elif col in PCT_COLS:
            fmt = "0.00"
        elif pd.api.types.is_integer_dtype(df[col]):
            fmt = "#,##0"
        if fmt:
            for row in range(2, len(df) + 2):
                ws.cell(row=row, column=i).number_format = fmt


def build_data_quality(source_dir: Path) -> pd.DataFrame:
    db = a.db_quality_checks()
    db.insert(0, "source", "Database")
    frames = [db]
    if (source_dir / "orders.csv").exists():
        summary, _ = validate_data.run_rules(source_dir)
        fv = summary.rename(columns={"rule": "check"})
        fv.insert(0, "source", "CSV files")
        frames.append(fv)
    return pd.concat(frames, ignore_index=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default="data", help="CSV folder to validate (default: data)")
    args = parser.parse_args()
    source_dir = (ROOT / args.source).resolve()
    REPORTS_DIR.mkdir(exist_ok=True)

    print("Pulling data from PostgreSQL views...")
    kpis = a.kpis()

    monthly = a.monthly_revenue().merge(
        a.monthly_growth()[["month", "growth_pct"]], on="month", how="left")
    monthly["month"] = monthly["month"].dt.strftime("%Y-%m")

    regions = a.region_revenue()
    categories = a.category_performance()
    products_all = a.read_view("v_top_products").sort_values("revenue", ascending=False)
    products_all["margin_pct"] = (100 * products_all["profit"] / products_all["revenue"]).round(2)
    top_products = products_all.head(25).reset_index(drop=True)
    customers = a.top_customers(25)
    customers["last_order_date"] = customers["last_order_date"].dt.strftime("%Y-%m-%d")

    print("Running data quality checks...")
    quality = build_data_quality(source_dir)

    summary = pd.DataFrame([
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Period covered", f"{monthly['month'].iloc[0]} to {monthly['month'].iloc[-1]}"),
        ("Total revenue (INR, completed orders)", kpis["total_revenue"]),
        ("Total orders", kpis["total_orders"]),
        ("Completed orders", kpis["completed_orders"]),
        ("Active customers", kpis["active_customers"]),
        ("Average order value (INR)", kpis["avg_order_value"]),
        ("Cancellation rate (%)", kpis["cancellation_rate_pct"]),
        ("Data quality checks failing", int((quality["status"] == "FAIL").sum())),
    ], columns=["Metric", "Value"])

    xlsx_path = REPORTS_DIR / "monthly_sales_report.xlsx"
    sheets = {
        "Summary": summary,
        "Monthly Revenue": monthly,
        "Regions": regions,
        "Categories": categories,
        "Top Products": top_products,
        "Top Customers": customers,
        "Data Quality": quality,
    }
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
        for name, df in sheets.items():
            style_sheet(writer.sheets[name], df)

        ws = writer.sheets["Summary"]
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 24
        for row in range(2, len(summary) + 2):
            v = ws.cell(row=row, column=2).value
            if isinstance(v, float):
                ws.cell(row=row, column=2).number_format = "#,##0.00"
            elif isinstance(v, int):
                ws.cell(row=row, column=2).number_format = "#,##0"
            ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")

        ws = writer.sheets["Monthly Revenue"]
        rev_col = list(monthly.columns).index("revenue") + 1
        chart = LineChart()
        chart.title = "Monthly revenue (INR)"
        chart.height, chart.width = 8, 18
        chart.add_data(Reference(ws, min_col=rev_col, min_row=1, max_row=len(monthly) + 1),
                       titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(monthly) + 1))
        ws.add_chart(chart, "G2")

    prod_path = REPORTS_DIR / "product_performance.csv"
    products_all.to_csv(prod_path, index=False)

    dq_path = REPORTS_DIR / "data_quality_report.csv"
    if (source_dir / "orders.csv").exists():
        file_summary, _ = validate_data.run_rules(source_dir)
        file_summary.to_csv(dq_path, index=False)
        dq_note = f"refreshed from {source_dir.name}/"
    else:
        dq_note = "kept existing file (no CSVs found to validate)"

    print("\nReports written to reports/:")
    print(f"  monthly_sales_report.xlsx   ({len(sheets)} sheets)")
    print(f"  product_performance.csv     ({len(products_all):,} products)")
    print(f"  data_quality_report.csv     ({dq_note})")
    print(f"\nFailing quality checks: {int((quality['status'] == 'FAIL').sum())}")


if __name__ == "__main__":
    main()
